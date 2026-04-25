import pdfplumber
import docx
from pptx import Presentation
import markdown

def read_pdf(path):
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text += page.extract_text() or ""
            text += "\n"
    return text


def read_docx(path):
    doc = docx.Document(path)
    return "\n".join([para.text for para in doc.paragraphs])


def read_pptx(path):
    prs = Presentation(path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text


def read_txt(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def read_md(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
        return markdown.markdown(content)     # HTML output


def process_document(path):
    if path.endswith(".pdf"):
        return read_pdf(path)

    elif path.endswith(".docx"):
        return read_docx(path)

    elif path.endswith(".pptx"):
        return read_pptx(path)

    elif path.endswith(".txt"):
        return read_txt(path)

    elif path.endswith(".md"):
        return read_md(path)

    return "Unsupported file format"
